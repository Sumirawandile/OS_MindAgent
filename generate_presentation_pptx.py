import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # Palette
    NAVY = RGBColor(0x1E, 0x3A, 0x8A)
    BLUE = RGBColor(0x25, 0x63, 0xEB)
    DARK = RGBColor(0x0F, 0x17, 0x2A)
    MUTED = RGBColor(0x64, 0x74, 0x8B)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)
    CARD_BG = RGBColor(0xEF, 0xF6, 0xFF)

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.color.rgb = LIGHT_BG
        return bg

    def add_header(slide, title_text, category_text="MINI PROJECT PRESENTATION"):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.0))
        tf = header_box.text_frame
        tf.word_wrap = True
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = BLUE
        
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

    def add_card(slide, left, top, width, height, title, items, bg_color=CARD_BG, border_color=BLUE):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_bottom = Inches(0.2)
        
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY
        p_title.space_after = Pt(10)
        
        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = f"• {item}"
            p_item.font.size = Pt(12)
            p_item.font.color.rgb = DARK
            p_item.space_after = Pt(6)

    # SLIDE 1: TITLE SLIDE
    slide1 = prs.slides.add_slide(blank_layout)
    add_bg(slide1)
    
    tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "OS-MindAgent: Automated Doubt Resolution System"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = NAVY
    p1.space_after = Pt(10)
    
    p2 = tf1.add_paragraph()
    p2.text = "Multi-Agent Architecture for Operating Systems Learning & Practice Evaluation"
    p2.font.size = Pt(18)
    p2.font.color.rgb = BLUE
    p2.space_after = Pt(30)
    
    p3 = tf1.add_paragraph()
    p3.text = "Submitted for Course: AGENTIC AI & AUTOMATION (B.Tech CSE)\nSymbiosis Institute of Technology (SIT), Nagpur Campus"
    p3.font.size = Pt(14)
    p3.font.color.rgb = DARK
    p3.space_after = Pt(15)
    
    p4 = tf1.add_paragraph()
    p4.text = "Under the Guidance of: Dr. Shreyas Rajendra Hole / Mr. Parag Naik"
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = MUTED

    # SLIDE 2: PRESENTATION FLOW PART 1
    slide2 = prs.slides.add_slide(blank_layout)
    add_bg(slide2)
    add_header(slide2, "1. Presentation Flow (Agenda - Part 1)")
    
    add_card(slide2, 0.8, 1.6, 5.6, 2.4, "01. Problem Statement", [
        "Abstract OS concepts are difficult to master.",
        "Static text explanations lack code & visual diagrams.",
        "Need interactive step-by-step doubt resolution."
    ])
    
    add_card(slide2, 6.8, 1.6, 5.6, 2.4, "02. Objectives & Research", [
        "Autonomous query classification into OS topics.",
        "Sub-1.2s LLM inference via Groq LPU API.",
        "Dynamic MCQ practice quiz generation."
    ])
    
    add_card(slide2, 0.8, 4.4, 5.6, 2.4, "03. Existing Solutions", [
        "Generic LLM chatbots lack syllabus context.",
        "No dynamic visual Gantt charts or state diagrams.",
        "Lack of topic-based automated scoring."
    ])

    add_card(slide2, 6.8, 4.4, 5.6, 2.4, "04. Compare & Contrast", [
        "Multi-agent Router-Solver vs single-prompt wrapper.",
        "Zero-downtime keyword fallback classifier.",
        "SQLite persistent memory vs stateless sessions."
    ])

    # SLIDE 3: PRESENTATION FLOW PART 2
    slide3 = prs.slides.add_slide(blank_layout)
    add_bg(slide3)
    add_header(slide3, "2. Presentation Flow (Agenda - Part 2)")

    add_card(slide3, 0.8, 1.6, 5.6, 2.4, "05. Problem Modeling & Routing", [
        "Router Agent categorizes doubts with confidence %.",
        "Recommends diagram tool based on intent.",
        "Decouples triage from answer synthesis."
    ])

    add_card(slide3, 6.8, 1.6, 5.6, 2.4, "06. Project Features", [
        "Tutor Agent generates Linux C code & viva formulas.",
        "Custom visual Gantt/TLB diagram rendering.",
        "ReportLab PDF export for revision sheets."
    ])

    add_card(slide3, 0.8, 4.4, 5.6, 2.4, "07. Results & Outcomes", [
        "Sub-1.2 second response latency.",
        "95%+ accuracy in OS topic classification.",
        "Safe MCQ evaluation with instant explanations."
    ])

    add_card(slide3, 6.8, 4.4, 5.6, 2.4, "08. Strengths & Future Scope", [
        "Modular multi-agent design pattern.",
        "High user engagement through visual tools.",
        "Future expansion: RAG textbook retrieval & MCP tools."
    ])

    # SLIDE 4: PROBLEM STATEMENT & MOTIVATION
    slide4 = prs.slides.add_slide(blank_layout)
    add_bg(slide4)
    add_header(slide4, "Problem Statement & Motivation")
    
    add_card(slide4, 0.8, 1.6, 11.6, 5.0, "Core Academic Challenges & Motivation", [
        "Operating Systems (OS) concepts like CPU Scheduling, Deadlocks, Paging, and File Systems are abstract.",
        "Students need more than static text: they need step-by-step reasoning, C code snippets, and visual state diagrams.",
        "Motivation: Implement Agentic AI design patterns (Router-Solver, Handoffs, Tools, Persistent Memory) learned in class.",
        "Goal: Create an interactive, high-speed Operating Systems AI Tutor with practice evaluation."
    ])

    # SLIDE 5: MULTI-AGENT SYSTEM ARCHITECTURE
    slide5 = prs.slides.add_slide(blank_layout)
    add_bg(slide5)
    add_header(slide5, "System Architecture & Agentic AI Workflow")

    add_card(slide5, 0.8, 1.6, 3.6, 5.0, "1. Router Agent", [
        "Classifies input into 5 OS topics.",
        "Calculates confidence score (e.g. 0.95).",
        "Recommends visual diagram tool."
    ])

    add_card(slide5, 4.8, 1.6, 3.6, 5.0, "2. Specialist Tutor", [
        "Powered by Groq LPU (Llama-3.3-70B).",
        "Generates 4-part structured markdown answer.",
        "Provides Linux C code & Viva tips."
    ])

    add_card(slide5, 8.8, 1.6, 3.6, 5.0, "3. Custom Tools & Memory", [
        "Mermaid Gantt & TLB Diagram Tool.",
        "Practice MCQ Quiz Engine.",
        "SQLite Database & PDF Report Exporter."
    ])

    # SLIDE 6: IMPLEMENTATION OF KEY FEATURES
    slide6 = prs.slides.add_slide(blank_layout)
    add_bg(slide6)
    add_header(slide6, "Implementation of Project Features")

    add_card(slide6, 0.8, 1.6, 5.6, 5.0, "Interactive Smart Doubt Resolver", [
        "Modern Gradio Web Interface (app.py).",
        "Preset sample queries for instant testing.",
        "Visual diagram toggle checkbox.",
        "High contrast royal blue aesthetic UI."
    ])

    add_card(slide6, 6.8, 1.6, 5.6, 5.0, "Practice Quiz & Report Generation", [
        "Topic-based 3-question MCQ evaluation.",
        "Robust string matching (safe against KeyErrors).",
        "Persistent logging of quiz scores in SQLite.",
        "Downloadable PDF revision summary sheets."
    ])

    # SLIDE 7: RESULTS & PERFORMANCE
    slide7 = prs.slides.add_slide(blank_layout)
    add_bg(slide7)
    add_header(slide7, "Results, Metrics & Evaluation")

    add_card(slide7, 0.8, 1.6, 5.6, 5.0, "Performance Metrics", [
        "Inference Speed: Sub-1.2 seconds via Groq LPU.",
        "Classification Accuracy: 95%+ across 5 OS topics.",
        "System Uptime: 100% via heuristic fallback system."
    ])

    add_card(slide7, 6.8, 1.6, 5.6, 5.0, "Student Outcomes", [
        "Faster doubt resolution before exams.",
        "Enhanced comprehension through Gantt/TLB diagrams.",
        "Effective self-assessment via topic practice quizzes."
    ])

    # SLIDE 8: CONCLUSION & Q/A
    slide8 = prs.slides.add_slide(blank_layout)
    add_bg(slide8)
    add_header(slide8, "Conclusion & Project Summary")

    add_card(slide8, 0.8, 1.6, 11.6, 5.0, "Summary & Viva Reference", [
        "OS-MindAgent successfully applies Agentic AI multi-agent architecture to Operating Systems education.",
        "Demonstrates Router-Solver pipeline, dynamic tool dispatch, SQLite persistent memory, and PDF reporting.",
        "GitHub Repository: Initialized & committed with secret protection.",
        "Thank You! Questions & Discussion."
    ])

    prs.save("OS_MindAgent_Mini_Project_Presentation.pptx")
    print("Presentation pptx generated successfully!")

if __name__ == "__main__":
    create_presentation()
